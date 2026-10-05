import os
import docx
import sys

input_file = r"C:\Users\Krrish Rebba\.gemini\antigravity-ide\brain\3f68e9e9-a079-46b4-afec-76b08ff902c9\Applied_Prototype_Report.md"
output_file = r"C:\Users\Krrish Rebba\OneDrive\ドキュメント\RM\Findings\ResearchPaper_Final.docx"

if not os.path.exists(input_file):
    print("Input file not found.")
    sys.exit(1)

with open(input_file, 'r', encoding='utf-8') as f:
    lines = f.readlines()

doc = docx.Document()

for line in lines:
    line = line.strip()
    if not line:
        continue
    
    if line.startswith('# '):
        doc.add_heading(line[2:], level=1)
    elif line.startswith('## '):
        doc.add_heading(line[3:], level=2)
    elif line.startswith('### '):
        doc.add_heading(line[4:], level=3)
    elif line.startswith('**'):
        p = doc.add_paragraph()
        run = p.add_run(line.replace('**', ''))
        run.bold = True
    elif line.startswith('- ') or line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. '):
        doc.add_paragraph(line, style='List Paragraph')
    else:
        doc.add_paragraph(line)

doc.save(output_file)
print(f"Successfully saved to {output_file}")
